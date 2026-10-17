"""Printable admission form for one student, as PDF or Excel.

The form carries the school header, the student's details and (when there is one)
the admission fee / books / uniform summary, with signature lines at the bottom.
"""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models.settings import SchoolSettings
from app.services import admission_fee


def _fmt_date(value):
    return value.strftime('%d %b %Y') if value else ''


def _context(student):
    settings = SchoolSettings.query.first()
    class_name = student.class_info.name if student.class_info else ''
    if student.section_info:
        class_name = '%s - %s' % (class_name, student.section_info.name)
    rows = [
        ('Admission No', student.admission_number or ''),
        ('Admission Date', _fmt_date(student.admission_date)),
        ("Student's Name", student.student_name),
        ('Gender', student.gender or ''),
        ('Date of Birth', _fmt_date(student.date_of_birth)),
        ('Class / Section', class_name),
        ('Roll No', student.roll_number),
        ("Guardian's Name", student.father_name),
        ('Guardian Relation', student.sponsor_type or ''),
        ('Guardian CNIC', student.sponsor_cnic or ''),
        ('Guardian Phone', student.guardian_phone),
        ('Address', student.address),
    ]
    fees = []
    info = admission_fee.summary(student)
    if info['applicable']:
        for line in info['lines']:
            fees.append((line['label'], line['qty'], admission_fee.to_rupees(line['unit']),
                         admission_fee.to_rupees(line['due']),
                         admission_fee.to_rupees(line['paid'])))
    return {
        'school': (settings.school_name if settings and settings.school_name else 'School'),
        'address': settings.address if settings else '',
        'phone': settings.phone if settings else '',
        'rows': rows, 'fees': fees, 'info': info,
        'total': admission_fee.to_rupees(info['total']),
        'received': admission_fee.to_rupees(info['received']),
        'balance': admission_fee.to_rupees(info['balance']),
    }


def build_pdf(student):
    ctx = _context(student)
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=15 * mm, bottomMargin=15 * mm,
                            title='Admission Form - %s' % student.student_name)
    base = getSampleStyleSheet()
    title = ParagraphStyle('t', parent=base['Title'], fontSize=18, spaceAfter=2)
    small = ParagraphStyle('s', parent=base['Normal'], fontSize=9, alignment=1,
                           textColor=colors.HexColor('#475569'))
    head = ParagraphStyle('h', parent=base['Heading3'], fontSize=11, spaceBefore=10,
                          spaceAfter=4, textColor=colors.HexColor('#0f172a'))
    cell = ParagraphStyle('c', parent=base['Normal'], fontSize=10, leading=13)
    story = [Paragraph(ctx['school'], title)]
    line = ' | '.join(x for x in (ctx['address'], ctx['phone']) if x)
    if line:
        story.append(Paragraph(line, small))
    story += [Spacer(1, 6), Paragraph('<b>ADMISSION FORM</b>', ParagraphStyle(
        'f', parent=base['Heading2'], alignment=1, fontSize=14)), Spacer(1, 4)]

    story.append(Paragraph('Student Details', head))
    data = [[Paragraph('<b>%s</b>' % k, cell), Paragraph(str(v), cell)] for k, v in ctx['rows']]
    table = Table(data, colWidths=[48 * mm, 126 * mm])
    table.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#94a3b8')),
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f1f5f9')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5)]))
    story.append(table)

    if ctx['fees']:
        story.append(Paragraph('Admission Fee, Books &amp; Uniform', head))
        fee_rows = [['Item', 'Qty', 'Price', 'Amount', 'Paid']]
        for label, qty, unit, due, paid in ctx['fees']:
            fee_rows.append([label, qty, '%.2f' % unit, '%.2f' % due, '%.2f' % paid])
        fee_rows.append(['Total', '', '', '%.2f' % ctx['total'], '%.2f' % ctx['received']])
        fee_rows.append(['Balance', '', '', '', '%.2f' % ctx['balance']])
        ft = Table(fee_rows, colWidths=[62 * mm, 16 * mm, 30 * mm, 33 * mm, 33 * mm])
        ft.setStyle(TableStyle([
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#94a3b8')),
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e2e8f0')),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTNAME', (0, -2), (-1, -1), 'Helvetica-Bold'),
            ('ALIGN', (1, 0), (-1, -1), 'RIGHT')]))
        story.append(ft)

    story += [Spacer(1, 38)]
    sign = Table([['_______________________', '_______________________', '_______________________'],
                  ["Guardian's Signature", 'Admission In-charge', 'Principal']],
                 colWidths=[58 * mm] * 3)
    sign.setStyle(TableStyle([('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                              ('FONTSIZE', (0, 1), (-1, 1), 9)]))
    story += [sign, Spacer(1, 10),
              Paragraph('Printed on %s' % date.today().strftime('%d %b %Y'), small)]
    doc.build(story)
    return out.getvalue()


def build_xlsx(student):
    ctx = _context(student)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Admission Form'
    ws.column_dimensions['A'].width = 24
    for col in 'BCDE':
        ws.column_dimensions[col].width = 18
    ws.merge_cells('A1:E1')
    ws['A1'] = ctx['school']
    ws['A1'].font = Font(bold=True, size=16)
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.merge_cells('A2:E2')
    ws['A2'] = ' | '.join(x for x in (ctx['address'], ctx['phone']) if x)
    ws['A2'].alignment = Alignment(horizontal='center')
    ws.merge_cells('A3:E3')
    ws['A3'] = 'ADMISSION FORM'
    ws['A3'].font = Font(bold=True, size=13)
    ws['A3'].alignment = Alignment(horizontal='center')
    row = 5
    label_fill = PatternFill('solid', fgColor='F1F5F9')
    for key, value in ctx['rows']:
        ws.cell(row=row, column=1, value=key).font = Font(bold=True)
        ws.cell(row=row, column=1).fill = label_fill
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        ws.cell(row=row, column=2, value=value)
        row += 1
    if ctx['fees']:
        row += 1
        ws.cell(row=row, column=1, value='Admission Fee, Books & Uniform').font = Font(bold=True)
        row += 1
        for col, name in enumerate(('Item', 'Qty', 'Price', 'Amount', 'Paid'), start=1):
            cell = ws.cell(row=row, column=col, value=name)
            cell.font = Font(bold=True)
            cell.fill = label_fill
        for label, qty, unit, due, paid in ctx['fees']:
            row += 1
            for col, value in enumerate((label, qty, unit, due, paid), start=1):
                ws.cell(row=row, column=col, value=value)
        row += 1
        for label, amount, col in (('Total', ctx['total'], 4), ('Received', ctx['received'], 5),
                                   ('Balance', ctx['balance'], 5)):
            ws.cell(row=row, column=1, value=label).font = Font(bold=True)
            ws.cell(row=row, column=col, value=amount).font = Font(bold=True)
            row += 1
    row += 2
    for col, name in ((1, "Guardian's Signature"), (3, 'Admission In-charge'), (5, 'Principal')):
        ws.cell(row=row, column=col, value='____________________')
        ws.cell(row=row + 1, column=col, value=name)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
