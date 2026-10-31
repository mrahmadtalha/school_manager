"""Teacher exports (CSV, Excel, PDF) and bulk import."""

import io
import csv
import json
import re
from flask import render_template, request, redirect, url_for, flash, send_file, Response, abort
from reportlab.lib.pagesizes import landscape, letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from app.database import db
from app.models import TeacherModel, SubjectModel, ClassModel
from app.routes import main
from app.services.photos import photo_path
from app.models.settings import get_custom_fields
from app.services import record_view
from app.services.custom_fields import serialize_custom_field_values
from app.services.data_import import (
    ImportFileError,
    ImportResult,
    RowChecker,
    build_template_workbook,
    collect_custom_values,
    is_example_row,
    norm_name,
    read_import_file,
    split_list,
    teacher_import_fields,
)

from app.routes.teachers.crud import _teacher_export_context


@main.route('/teachers/export/csv')
def export_teachers_csv():
    teachers, specs, extras = _teacher_export_context()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(record_view.headers(specs))
    for row in record_view.build_rows(teachers, specs, extras):
        writer.writerow(row)
    output.seek(0)
    return Response(output.getvalue(), mimetype="text/csv", headers={"Content-Disposition": "attachment;filename=teachers_report.csv"})


TEACHER_EXAMPLE_KEYS = ('teacher_name', 'qualification', 'cnic')


@main.route('/teachers/template/excel')
def teachers_template_excel():
    """Import template: every current field, one example row, mandatory marked."""
    classes = ClassModel.query.order_by(ClassModel.name).all()
    subjects = [row[0] for row in db.session.query(SubjectModel.name)
                .distinct().order_by(SubjectModel.name).all()]
    specs = teacher_import_fields(get_custom_fields('teacher'))
    example = {spec.key: spec.example for spec in specs}
    if classes:
        example['assigned_classes'] = classes[0].name
    if subjects:
        example['assigned_subjects'] = subjects[0]

    intro = [
        'How to fill in the Teachers sheet:',
        '1. Keep the column headings in row 1 exactly as they are. Do not rename or delete them.',
        '2. The yellow row 2 is only an EXAMPLE. Type over it (or delete it). If it is left unchanged it is ignored.',
        '3. Add one teacher per row, starting from row 2. Red columns must always be filled in.',
        '4. Teacher ID is optional: leave it blank and the next free ID (T001, T002 ...) is given automatically.',
        '5. Teachers are never overwritten. A row is skipped (and reported) if the Teacher ID or CNIC is already '
        'registered, or the same teacher (same name and contact number, or same name, joining date and '
        'qualification) already exists.',
        '6. Assigned Classes / Assigned Subjects must already exist in the system; separate several with commas.',
        '7. Save the file as .xlsx (or .csv) and upload it with Import Excel. You will see exactly which rows '
        'were skipped and why.',
        'Your classes: ' + (', '.join(c.name for c in classes) or 'none yet'),
        'Your subjects: ' + (', '.join(subjects) or 'none yet'),
    ]
    output = build_template_workbook(
        'Teachers', specs, example, 'Instructions', intro,
        lists={'Classes': [c.name for c in classes]},
        dropdowns={
            'gender': ('list', ['Male', 'Female', 'Other']),
            'salary_type': ('list', ['monthly', 'hourly']),
            'designation': ('list', ['Principal', 'Vice Principal', 'Head Teacher',
                                     'Senior Teacher', 'Teacher', 'Admin Staff', 'Other']),
        })
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='teachers_import_template.xlsx')


@main.route('/teachers/export/excel')
def export_teachers_excel():
    teachers, specs, extras = _teacher_export_context()
    headers = record_view.headers(specs)
    rows = [[record_view.excel_value(spec, teacher, extras) for spec in specs]
            for teacher in teachers]

    wb = Workbook()
    ws = wb.active
    ws.title = 'Teachers'
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='2D3748')
    for row_index, row in enumerate(rows, 2):
        for col_index, value in enumerate(row, 1):
            ws.cell(row=row_index, column=col_index, value=value)
    ws.freeze_panes = 'A2'
    if rows:
        ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{len(rows) + 1}'
    for col in range(1, len(headers) + 1):
        width = max([len(str(headers[col - 1]))] +
                    [len(str(row[col - 1])) for row in rows] + [6])
        ws.column_dimensions[get_column_letter(col)].width = min(width + 2, 32)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name='teachers_report.xlsx')


@main.route('/teachers/export/pdf')
def export_teachers_pdf():
    teachers, specs, extras = _teacher_export_context()
    headers = record_view.headers(specs)
    rows = record_view.build_rows(teachers, specs, extras)

    output = io.BytesIO()
    wide = len(specs) > 7
    doc = SimpleDocTemplate(output, pagesize=landscape(letter) if wide else letter,
                            rightMargin=24, leftMargin=24, topMargin=24, bottomMargin=24)
    story = []
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontSize=16, textColor=colors.HexColor('#1a202c'), spaceAfter=12, alignment=1)

    story.append(Paragraph("School Management System — Teacher Report", title_style))
    story.append(Spacer(1, 6))

    if not specs:
        story.append(Paragraph("No columns selected.", styles['Normal']))
    else:
        font_size = 8 if len(specs) <= 10 else 6
        table = Table([headers] + rows, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2d3748')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), font_size),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#f7fafc')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e0')),
            ('FONTSIZE', (0, 0), (-1, -1), font_size),
        ]))
        story.append(table)
    doc.build(story)
    output.seek(0)
    return send_file(output, mimetype='application/pdf', as_attachment=True, download_name='teachers_report.pdf')


def _digits_only(value):
    return re.sub(r'\D', '', value or '')


@main.route('/teachers/import', methods=['POST'])
def import_teachers():
    """Bulk-import teachers.

    Rules: mandatory fields must be present; nothing already registered is ever
    overwritten (duplicate Teacher ID, duplicate CNIC, or the same person
    registered before are skipped); a blank Teacher ID is assigned automatically
    (next free T-number).  Every skipped row is listed with its reason on the
    result page.
    """
    file = request.files.get('import_file')
    if not file or not file.filename:
        flash('No file selected.', 'warning')
        return redirect(url_for('main.teachers_list'))

    specs = teacher_import_fields(get_custom_fields('teacher'))
    try:
        parsed = read_import_file(file, specs)
    except ImportFileError as error:
        flash(str(error), 'danger')
        return redirect(url_for('main.teachers_list'))

    result = ImportResult('teacher')
    result.total_rows = len(parsed.rows)
    result.ignored_columns = parsed.ignored_columns

    class_lookup = {c.name.strip().lower(): c.name for c in ClassModel.query.all()}
    subject_lookup = {row[0].strip().lower(): row[0]
                      for row in db.session.query(SubjectModel.name).distinct().all() if row[0]}

    # What is already registered (archived teachers count too).
    used_ids = {}       # lower teacher id -> 'registered' | file row number
    used_cnic = {}      # cnic -> 'registered' | file row number
    people = {}         # normalised name -> [dict(contact, joining, qualification, source, active)]
    for row in TeacherModel.query.with_entities(
            TeacherModel.teacher_id_str, TeacherModel.teacher_name, TeacherModel.cnic,
            TeacherModel.contact_number, TeacherModel.joining_date,
            TeacherModel.qualification, TeacherModel.is_active).all():
        used_ids[(row.teacher_id_str or '').strip().lower()] = 'registered'
        if row.cnic:
            used_cnic[row.cnic.strip()] = 'registered'
        joined = row.joining_date.date() if hasattr(row.joining_date, 'date') else row.joining_date
        people.setdefault(norm_name(row.teacher_name), []).append({
            'contact': _digits_only(row.contact_number), 'joining': joined,
            'qualification': norm_name(row.qualification), 'source': 'registered',
            'active': bool(row.is_active)})

    example = {spec.key: spec.example for spec in specs}
    pending = []

    # ---- pass 1: validate every row, detect duplicates ----------------------
    for row_no, record in parsed.rows:
        if is_example_row(record, example, TEACHER_EXAMPLE_KEYS):
            result.example_rows += 1
            continue

        chk = RowChecker(record, specs)
        name = chk.text('teacher_name', required=True, max_len=100)
        joining = chk.date('joining_date', required=True)
        qualification = chk.text('qualification', required=True, max_len=100)
        teacher_id = chk.text('teacher_id_str', max_len=50)
        if ' ' in teacher_id:
            chk.add_problem(f'Teacher ID "{teacher_id}" must not contain spaces')
        designation = chk.text('designation', max_len=80)
        gender = chk.choice('gender', ('Male', 'Female', 'Other'),
                            aliases={'m': 'Male', 'f': 'Female'})
        salary_type = chk.choice('salary_type', ('monthly', 'hourly'), default='monthly')
        monthly_salary = chk.number('monthly_salary')
        hourly_rate = chk.number('hourly_rate')
        dob = chk.date('date_of_birth')
        cnic = chk.cnic('cnic')
        address = chk.long_text('address')
        contact = chk.phone('contact_number')
        emergency = chk.phone('emergency_contact_number')
        experience = chk.number('previous_experience_years')
        previous_salary = chk.number('previous_salary')
        custom_values = collect_custom_values(chk, specs)

        classes_chosen, unknown = [], []
        for item in split_list(chk.raw('assigned_classes')):
            found = class_lookup.get(item.lower())
            (classes_chosen if found else unknown).append(found or item)
        if unknown:
            chk.add_problem('Assigned Classes not found: ' + ', '.join(unknown))
        subjects_chosen, unknown = [], []
        for item in split_list(chk.raw('assigned_subjects')):
            found = subject_lookup.get(item.lower())
            (subjects_chosen if found else unknown).append(found or item)
        if unknown:
            chk.add_problem('Assigned Subjects not found: ' + ', '.join(unknown))

        reason = chk.reason()
        if reason:
            result.skip(row_no, name, reason, 'missing' if chk.missing else 'invalid')
            continue

        # duplicates -------------------------------------------------------
        duplicates = []
        if teacher_id and teacher_id.lower() in used_ids:
            source = used_ids[teacher_id.lower()]
            duplicates.append(
                f'Teacher ID {teacher_id} is repeated in this file (already used in row {source})'
                if isinstance(source, int) else f'Teacher ID {teacher_id} is already registered')
        if cnic and cnic in used_cnic:
            source = used_cnic[cnic]
            duplicates.append(
                f'CNIC {cnic} is repeated in this file (already used in row {source})'
                if isinstance(source, int) else f'CNIC {cnic} is already registered')
        if not duplicates:
            for other in people.get(norm_name(name), []):
                same_contact = bool(contact) and _digits_only(contact) == other['contact']
                same_profile = (joining == other['joining']
                                and norm_name(qualification) == other['qualification'])
                if same_contact or same_profile:
                    if isinstance(other['source'], int):
                        duplicates.append(f'Looks like a duplicate of row {other["source"]} in this file '
                                          '(same name and contact number / joining date)')
                    else:
                        duplicates.append(
                            f'A teacher named {name} with the same contact number or joining date is '
                            'already registered' + ('' if other['active'] else ' (archived - restore instead)'))
                    break
        if duplicates:
            result.skip(row_no, name, '; '.join(duplicates), 'duplicate')
            continue

        if teacher_id:
            used_ids[teacher_id.lower()] = row_no
        if cnic:
            used_cnic[cnic] = row_no
        people.setdefault(norm_name(name), []).append({
            'contact': _digits_only(contact), 'joining': joining,
            'qualification': norm_name(qualification), 'source': row_no, 'active': True})
        pending.append({
            'row': row_no, 'name': name, 'joining': joining, 'qualification': qualification,
            'teacher_id': teacher_id, 'designation': designation or None, 'gender': gender,
            'salary_type': salary_type, 'monthly_salary': monthly_salary,
            'hourly_rate': hourly_rate, 'dob': dob, 'cnic': cnic, 'address': address or None,
            'contact': contact or None, 'emergency': emergency or None,
            'classes': classes_chosen, 'subjects': subjects_chosen,
            'experience': experience, 'previous_salary': previous_salary,
            'custom': custom_values,
        })

    # ---- pass 2: assign automatic Teacher IDs and save ----------------------
    def _next_auto_id():
        highest = 0
        for existing in used_ids:
            match = re.fullmatch(r't(\d+)', existing)
            if match:
                highest = max(highest, int(match.group(1)))
        number = highest + 1
        while ('t%03d' % number) in used_ids:
            number += 1
        return 'T%03d' % number

    try:
        for item in pending:
            auto_assigned = not item['teacher_id']
            if auto_assigned:
                item['teacher_id'] = _next_auto_id()
                used_ids[item['teacher_id'].lower()] = item['row']

            if item['salary_type'] == 'hourly':
                hourly = item['hourly_rate'] if item['hourly_rate'] is not None else (item['monthly_salary'] or 0.0)
                monthly, salary = 0.0, hourly
                hourly_value = hourly
            else:
                monthly = item['monthly_salary'] if item['monthly_salary'] is not None else (item['hourly_rate'] or 0.0)
                hourly_value, salary = 0.0, monthly
            if not salary:
                result.warn(item['row'], item['name'],
                            'No salary was given, so the salary is 0. Edit the teacher to set it.')

            classes = item['classes']
            teacher = TeacherModel(
                teacher_id_str=item['teacher_id'],
                teacher_name=item['name'],
                joining_date=item['joining'],
                qualification=item['qualification'],
                salary=salary,
                salary_type=item['salary_type'],
                monthly_salary=monthly,
                hourly_rate=hourly_value,
                assigned_class=', '.join(classes) if classes else None,
                assigned_classes=json.dumps(classes) if classes else None,
                assigned_subjects=json.dumps(item['subjects']) if item['subjects'] else None,
                cnic=item['cnic'],
                designation=item['designation'],
                gender=item['gender'],
                address=item['address'],
                contact_number=item['contact'],
                emergency_contact_number=item['emergency'],
                previous_experience_years=item['experience'],
                previous_salary=item['previous_salary'],
                date_of_birth=item['dob'],
                custom_fields_data=serialize_custom_field_values(item['custom']),
                is_active=True,
            )
            db.session.add(teacher)
            result.added.append((item['row'], f"{item['name']} (ID {item['teacher_id']})" +
                                 (' - ID assigned automatically' if auto_assigned else '')))
        db.session.commit()
    except Exception as error:
        db.session.rollback()
        flash(f'Import cancelled and nothing was saved: {error}', 'danger')
        return redirect(url_for('main.teachers_list'))

    result.skipped.sort(key=lambda item: item['row'])
    return render_template('import_result.html', result=result, entity_label='Teacher',
                           entity_plural='Teachers',
                           back_url=url_for('main.teachers_list'),
                           template_url=url_for('main.teachers_template_excel'))


@main.route('/teachers/restore-all', methods=['POST'])
def restore_all_teachers():
    count = TeacherModel.query.filter_by(is_active=False).update({'is_active': True})
    db.session.commit()
    flash(f'{count} teacher(s) restored successfully.', 'success')
    return redirect(url_for('main.archived_teachers_list'))


@main.route('/teachers/<int:id>/photo')
def teacher_photo(id):
    teacher = TeacherModel.query.get_or_404(id)
    if not teacher.photo_filename:
        abort(404)
    path = photo_path('teacher', teacher.photo_filename)
    if not path:
        abort(404)
    return send_file(path, mimetype='image/jpeg')
