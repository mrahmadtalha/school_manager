"""Student list exports (CSV, Excel, PDF) and the import template."""

import io
import csv
from flask import request, send_file, Response
from reportlab.lib.pagesizes import landscape, letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from app.models import StudentModel, ClassModel
from app.routes import main
from app.services import record_view
from app.services import admission_fee
from app.models.settings import get_custom_fields
from app.services.data_import import build_template_workbook, student_import_fields

from app.routes.students.helpers import (
    _apply_student_filters,
    _dues_balances,
    _import_admission_enabled,
    _student_filters,
)


def _student_export_context():
    """Students to export for the current filters + the chosen columns.

    ``cols`` follows the column picker: absent or ``all`` exports every field,
    otherwise only the requested (and recognised) columns.
    """
    filters = _student_filters()
    dues_ids = ([sid for sid, bal in _dues_balances().items() if bal > 0]
                if filters['has_dues'] else None)
    students = (_apply_student_filters(StudentModel.query, filters,
                                       dues_ids=dues_ids)
                .order_by(StudentModel.student_name).all())
    dues_map = _dues_balances([s.id for s in students]) if students else {}
    specs = record_view.export_specs('student', request.args.get('cols'))
    return students, specs, {
        'dues_map': dues_map,
        'admission_map': admission_fee.student_map([s.id for s in students])}


@main.route('/students/export/csv')
def export_students_csv():
    students, specs, extras = _student_export_context()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(record_view.headers(specs))
    for row in record_view.build_rows(students, specs, extras):
        writer.writerow(row)

    output.seek(0)
    return Response(output.getvalue(), mimetype="text/csv", headers={"Content-Disposition": "attachment;filename=students_report.csv"})


@main.route('/students/template/excel')
def students_template_excel():
    """Import template: every current field, one example row, mandatory marked."""
    classes = ClassModel.query.order_by(ClassModel.name).all()
    specs = student_import_fields(get_custom_fields('student'),
                                 admission=_import_admission_enabled())
    example = {spec.key: spec.example for spec in specs}
    if classes:
        example['class_name'] = classes[0].name
        first_sections = sorted(classes[0].sections, key=lambda sec: sec.name)
        example['section'] = first_sections[0].name if first_sections else ''

    class_lines = '; '.join(
        f"{c.name} (sections: {', '.join(sec.name for sec in sorted(c.sections, key=lambda x: x.name))})"
        if c.sections else c.name
        for c in classes) or 'No classes created yet - add classes first.'
    intro = [
        'How to fill in the Students sheet:',
        '1. Keep the column headings in row 1 exactly as they are. Do not rename or delete them.',
        '2. The yellow row 2 is only an EXAMPLE. Type over it (or delete it). If it is left unchanged it is ignored.',
        '3. Add one student per row, starting from row 2. Red columns must always be filled in.',
        '4. Roll Number is optional: leave it blank and the next free number in the class is given automatically.',
        '5. Students are never overwritten. A row is skipped (and reported) if the Roll Number is already used '
        'in that class, the Admission No already exists, or a student with the same name, father name and class '
        'is already registered.',
        '6. Save the file as .xlsx (or .csv) and upload it with Import Excel. You will see exactly which rows '
        'were skipped and why.',
        'Your classes: ' + class_lines,
    ]
    output = build_template_workbook(
        'Students', specs, example, 'Instructions', intro,
        lists={'Classes': [c.name for c in classes]},
        dropdowns={
            'class_name': ('range', 'Classes'),
            'gender': ('list', ['Male', 'Female', 'Other']),
            'sponsor_type': ('list', ['Father', 'Guardian']),
            'discount_type': ('list', ['percentage', 'fixed']),
        })
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='students_import_template.xlsx')


@main.route('/students/export/excel')
def export_students_excel():
    students, specs, extras = _student_export_context()
    headers = record_view.headers(specs)
    rows = [[record_view.excel_value(spec, student, extras) for spec in specs]
            for student in students]

    wb = Workbook()
    ws = wb.active
    ws.title = 'Students'
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='2D3748')
    for row_index, row in enumerate(rows, 2):
        for col_index, value in enumerate(row, 1):
            ws.cell(row=row_index, column=col_index, value=value)
    ws.freeze_panes = 'A2'
    if rows:
        ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{len(rows) + 1}'
    for col, _ in enumerate(headers, 1):
        width = max([len(str(headers[col - 1]))] +
                    [len(str(row[col - 1])) for row in rows] + [6])
        ws.column_dimensions[get_column_letter(col)].width = min(width + 2, 32)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name='students_report.xlsx')


@main.route('/students/export/pdf')
def export_students_pdf():
    students, specs, extras = _student_export_context()
    headers = record_view.headers(specs)
    rows = record_view.build_rows(students, specs, extras)

    output = io.BytesIO()
    # All-fields exports are wide, so the page adapts to the column count.
    wide = len(specs) > 7
    pagesize = landscape(letter) if wide else letter
    doc = SimpleDocTemplate(output, pagesize=pagesize, rightMargin=24, leftMargin=24,
                            topMargin=24, bottomMargin=24)
    story = []
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontSize=16, textColor=colors.HexColor('#1a202c'), spaceAfter=12, alignment=1)

    story.append(Paragraph("School Management System — Student Report", title_style))
    story.append(Spacer(1, 6))

    if not specs:
        story.append(Paragraph("No columns selected.", styles['Normal']))
    else:
        table_data = [headers] + rows
        font_size = 8 if len(specs) <= 10 else 6
        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2d3748')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), font_size),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#f7fafc')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e0')),
            ('FONTSIZE', (0, 0), (-1, -1), font_size),
        ]))
        story.append(table)
    doc.build(story)
    output.seek(0)
    return send_file(output, mimetype='application/pdf', as_attachment=True, download_name='students_report.pdf')
