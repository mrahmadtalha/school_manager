"""Automatic Admission No (settings switch) and the admission form download."""
import io

from openpyxl import load_workbook

from app.database import db
from app.models import ClassModel
from app.services import admission_numbers
from tests.test_admission_fee import (  # noqa: F401  (fixtures are used by name)
    accountant_client, admin_client, teacher_client, _add_student, _student)


def _set_auto(client, on):
    data = {'action': 'save_general_info', 'school_name': 'Test School'}
    if on:
        data['auto_admission_number'] = 'on'
    resp = client.post('/settings', data=data)
    assert resp.status_code == 302


def test_auto_admission_number_is_off_until_switched_on(admin_client, seed):
    assert admission_numbers.is_auto() is False
    _add_student(admin_client, seed, 3001, 'Manual One', admission_number='REG-77')
    assert _student(3001).admission_number == 'REG-77'
    _add_student(admin_client, seed, 3002, 'Manual Two')
    assert _student(3002).admission_number is None


def test_switch_on_numbers_new_students_one_after_another(admin_client, seed):
    _set_auto(admin_client, True)
    assert admission_numbers.is_auto() is True
    page = admin_client.get('/settings').get_data(as_text=True)
    assert 'autoAdmissionNumber' in page and 'checked' in page
    _add_student(admin_client, seed, 3003, 'Auto One')
    _add_student(admin_client, seed, 3004, 'Auto Two')
    assert _student(3003).admission_number == 'ADM-0001'
    assert _student(3004).admission_number == 'ADM-0002'
    # the Add Student form shows the next number and the box is read-only
    html = admin_client.get('/students').get_data(as_text=True)
    assert 'Auto: ADM-0003' in html


def test_numbers_already_on_record_are_never_changed_or_reused(admin_client, seed):
    _add_student(admin_client, seed, 3005, 'Old Typed', admission_number='ADM-0010')
    _set_auto(admin_client, True)
    _add_student(admin_client, seed, 3006, 'New Auto')
    assert _student(3005).admission_number == 'ADM-0010'
    assert _student(3006).admission_number == 'ADM-0011'


def test_switch_off_goes_back_to_typing_by_hand(admin_client, seed):
    _set_auto(admin_client, True)
    _set_auto(admin_client, False)
    assert admission_numbers.is_auto() is False
    _add_student(admin_client, seed, 3007, 'Typed Again', admission_number='X-1')
    assert _student(3007).admission_number == 'X-1'


def test_admission_form_downloads_as_pdf_and_excel(admin_client, seed):
    class_obj = db.session.get(ClassModel, seed['class_id'])
    class_obj.admission_fee, class_obj.books_fee = 2000, 3000
    class_obj.uniform_fee, class_obj.uniform_sets = 1500, 2
    db.session.commit()
    _add_student(admin_client, seed, 3008, 'Form Kid',
                 admission_received_now='1', admission_received_amount='6000')
    student = _student(3008)
    pdf = admin_client.get('/students/%d/admission-form/pdf' % student.id)
    assert pdf.status_code == 200 and pdf.mimetype == 'application/pdf'
    assert pdf.data.startswith(b'%PDF') and 'attachment' in pdf.headers['Content-Disposition']
    xls = admin_client.get('/students/%d/admission-form/xlsx' % student.id)
    assert xls.status_code == 200 and xls.data[:2] == b'PK'
    sheet = load_workbook(io.BytesIO(xls.data)).active
    text = ' '.join(str(c.value) for row in sheet.iter_rows() for c in row if c.value is not None)
    for expected in ('Form Kid', 'ADMISSION FORM', 'Books', 'Uniform', 'Balance'):
        assert expected in text


def test_admission_form_for_a_student_without_fees_and_bad_format(admin_client, seed):
    _add_student(admin_client, seed, 3009, 'Plain Form')
    student = _student(3009)
    assert admin_client.get('/students/%d/admission-form/pdf' % student.id).status_code == 200
    assert admin_client.get('/students/%d/admission-form/doc' % student.id).status_code == 404
    assert admin_client.get('/students/999999/admission-form/pdf').status_code == 404


# ---- user-defined pattern ---------------------------------------------------

def _pattern(client, prefix, digits='4', start='1', on=True):
    data = {'action': 'save_general_info', 'school_name': 'Test School',
            'admission_number_prefix': prefix, 'admission_number_digits': digits,
            'admission_number_start': start}
    if on:
        data['auto_admission_number'] = 'on'
    assert client.post('/settings', data=data).status_code == 302


def test_school_pattern_alberoni_style(admin_client, seed):
    _pattern(admin_client, 'alberoni-26-')
    for roll in (3101, 3102, 3103):
        _add_student(admin_client, seed, roll, 'Pattern %d' % roll)
    assert [_student(r).admission_number for r in (3101, 3102, 3103)] == [
        'alberoni-26-0001', 'alberoni-26-0002', 'alberoni-26-0003']


def test_year_token_digits_and_start(admin_client, seed):
    from datetime import date
    _pattern(admin_client, 'SCH-{YY}/', digits='6', start='100')
    _add_student(admin_client, seed, 3104, 'Year Kid')
    yy = '%02d' % (date.today().year % 100)
    assert _student(3104).admission_number == 'SCH-%s/000100' % yy
    assert admission_numbers.next_number() == 'SCH-%s/000101' % yy


def test_new_prefix_starts_again_and_old_numbers_stay(admin_client, seed):
    _pattern(admin_client, 'alberoni-25-')
    _add_student(admin_client, seed, 3105, 'Last Year')
    _pattern(admin_client, 'alberoni-26-')
    _add_student(admin_client, seed, 3106, 'This Year')
    assert _student(3105).admission_number == 'alberoni-25-0001'
    assert _student(3106).admission_number == 'alberoni-26-0001'


def test_a_bad_pattern_is_refused_and_the_old_one_kept(admin_client, seed):
    _pattern(admin_client, 'GOOD-')
    _pattern(admin_client, 'bad-{nope}-')
    assert admission_numbers.get_config()['prefix'] == 'GOOD-'
    for prefix, digits in (('x' * 40, '4'), ('ok-', '0'), ('ok-', '9'), ('a;b', '4')):
        try:
            admission_numbers.save_config(prefix, digits, '1')
        except ValueError:
            continue
        raise AssertionError('accepted %r %r' % (prefix, digits))


def test_fee_reminders_month_is_a_dropdown_not_typed(admin_client, seed):
    html = admin_client.get('/fees/reminders').get_data(as_text=True)
    assert '<select class="form-select form-select-sm" name="month_year"' in html
    assert 'placeholder="e.g. September 2026"' not in html
