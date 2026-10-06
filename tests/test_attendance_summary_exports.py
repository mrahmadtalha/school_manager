"""Tests for whole-school attendance summary, holidays and rebuilt exports."""

import csv
import io
from datetime import date

from app.database import db
from app.models import (
    AttendanceModel, ClassModel, SchoolSettings, StudentModel, TeacherModel,
    get_holiday_ranges, set_holiday_ranges,
)
from app.services.attendance_service import (
    build_school_attendance_summary,
    build_student_attendance_summary,
    build_teacher_attendance_summary,
)
from app.services.teacher_payroll import (
    get_school_working_days, is_school_day,
)

# Monday..Friday inside one clean week; 2026-09-21 is a Monday.
MONDAY = date(2026, 9, 21)
LATE_TIME = '08:47'
LATE_MINUTES = 17
TUESDAY = date(2026, 9, 22)
WEDNESDAY = date(2026, 9, 23)
THURSDAY = date(2026, 9, 24)
FRIDAY = date(2026, 9, 25)
SATURDAY = date(2026, 9, 26)
SUNDAY = date(2026, 9, 27)


def _settings(**overrides):
    settings = SchoolSettings.query.first()
    if settings is None:
        settings = SchoolSettings()
        db.session.add(settings)
        db.session.flush()
    settings.school_start_time = '08:30'
    settings.attendance_grace_minutes = 0
    settings.saturday_off = True
    settings.sunday_off = True
    settings.custom_off_days = ''
    for key, value in overrides.items():
        setattr(settings, key, value)
    return settings


def _clear_holidays():
    set_holiday_ranges([])


def _add_student(class_obj, roll, name):
    student = StudentModel(
        roll_number=roll, student_name=name, father_name='Father',
        guardian_phone='03001112223', address='Street', class_id=class_obj.id,
        monthly_fee=1000.0,
    )
    db.session.add(student)
    db.session.flush()
    return student


def _add_attendance(student, day, status, class_id, late_minutes=None):
    db.session.add(AttendanceModel(
        target_type='student', target_id=student.id, class_id=class_id,
        date=day, status=status, late_minutes=late_minutes,
    ))


# ==========================================
# WORKING DAYS / HOLIDAYS
# ==========================================

def test_weekends_are_not_working_days(app):
    with app.app_context():
        settings = _settings()
        _clear_holidays()
        db.session.commit()
        assert not is_school_day(SATURDAY, settings)
        assert not is_school_day(SUNDAY, settings)
        working = get_school_working_days(MONDAY, SUNDAY, settings)
        assert working == [MONDAY, TUESDAY, WEDNESDAY, THURSDAY, FRIDAY]


def test_separate_saturday_and_sunday_toggles(app):
    with app.app_context():
        settings = _settings(saturday_off=False, sunday_off=True)
        db.session.commit()
        assert is_school_day(SATURDAY, settings)
        assert not is_school_day(SUNDAY, settings)

        settings.saturday_off = True
        settings.sunday_off = False
        db.session.commit()
        assert not is_school_day(SATURDAY, settings)
        assert is_school_day(SUNDAY, settings)


def test_holiday_range_days_are_not_working_days(app):
    with app.app_context():
        settings = _settings()
        set_holiday_ranges([
            {'label': 'Autumn Break', 'start': WEDNESDAY.isoformat(),
             'end': THURSDAY.isoformat()},
        ])
        working = get_school_working_days(MONDAY, FRIDAY, settings)
        assert working == [MONDAY, TUESDAY, FRIDAY]
        assert not is_school_day(WEDNESDAY, settings)
        assert not is_school_day(THURSDAY, settings)


def test_holiday_range_settings_round_trip(app):
    with app.app_context():
        set_holiday_ranges([
            {'label': 'Summer Vacations', 'start': '2026-06-01', 'end': '2026-08-15'},
            {'label': 'Eid Holidays', 'start': '2026-03-19', 'end': '2026-03-21'},
        ])
        stored = get_holiday_ranges()
        assert stored == [
            {'label': 'Summer Vacations', 'start': '2026-06-01', 'end': '2026-08-15'},
            {'label': 'Eid Holidays', 'start': '2026-03-19', 'end': '2026-03-21'},
        ]
        # Invalid rows (missing or reversed dates) are dropped, not stored.
        set_holiday_ranges([
            {'label': 'Bad', 'start': '2026-05-10', 'end': '2026-05-01'},
            {'label': 'Incomplete', 'start': '', 'end': '2026-05-01'},
            {'label': 'Good', 'start': '2026-05-02', 'end': '2026-05-03'},
        ])
        assert get_holiday_ranges() == [
            {'label': 'Good', 'start': '2026-05-02', 'end': '2026-05-03'}]


def test_settings_page_saves_saturday_sunday_and_holidays(admin_client, app):
    page = admin_client.get('/settings?tab=academic').get_data(as_text=True)
    assert 'name="saturday_off"' in page
    assert 'name="sunday_off"' in page
    assert 'Holidays / Vacations' in page
    assert 'name="holiday_start[]"' in page

    response = admin_client.post('/settings', data={
        'action': 'save_academic_settings',
        'school_start_time': '08:00',
        'school_end_time': '14:30',
        'saturday_off': 'on',
        'custom_off_days': 'Friday',
        'holiday_label[]': ['Winter Break'],
        'holiday_start[]': ['2026-12-20'],
        'holiday_end[]': ['2026-12-31'],
    })
    assert response.status_code == 302
    assert response.headers['Location'].endswith('/settings?tab=academic')

    with app.app_context():
        settings = SchoolSettings.query.first()
        assert settings.saturday_off is True
        assert settings.sunday_off is False
        assert settings.custom_off_days == 'Friday'
        assert get_holiday_ranges() == [
            {'label': 'Winter Break', 'start': '2026-12-20', 'end': '2026-12-31'}]


# ==========================================
# SUMMARY: HOLIDAYS EXCLUDED
# ==========================================

def test_summary_ignores_attendance_recorded_on_weekend_and_holiday(app, seed):
    with app.app_context():
        settings = _settings()
        set_holiday_ranges([
            {'label': 'Break', 'start': WEDNESDAY.isoformat(),
             'end': WEDNESDAY.isoformat()},
        ])
        class_obj = db.session.get(ClassModel, seed['class_id'])
        student = db.session.get(StudentModel, seed['student_id'])

        _add_attendance(student, MONDAY, 'Present', class_obj.id)
        # These two must never surface as attendance days.
        _add_attendance(student, SATURDAY, 'Present', class_obj.id)
        _add_attendance(student, WEDNESDAY, 'Absent', class_obj.id)
        db.session.commit()

        summary = build_student_attendance_summary(class_obj.id, MONDAY, SUNDAY)
        assert summary['dates_list'] == [MONDAY]
        assert SATURDAY not in summary['dates_list']
        assert WEDNESDAY not in summary['dates_list']

        row = summary['matrix_data'][0]
        assert row['p_count'] == 1
        assert row['a_count'] == 0
        assert row['percentage'] == 100.0

        teacher_summary = build_teacher_attendance_summary(MONDAY, SUNDAY)
        assert teacher_summary['dates_list'] == []


def test_summary_falls_back_to_all_working_days_without_records(app, seed):
    with app.app_context():
        _settings()
        _clear_holidays()
        db.session.commit()
        class_obj = db.session.get(ClassModel, seed['class_id'])
        summary = build_student_attendance_summary(class_obj.id, MONDAY, SUNDAY)
        assert summary['dates_list'] == [MONDAY, TUESDAY, WEDNESDAY, THURSDAY, FRIDAY]


# ==========================================
# WHOLE-SCHOOL SUMMARY
# ==========================================

def test_whole_school_summary_covers_every_class(app, seed):
    with app.app_context():
        _settings()
        _clear_holidays()
        second_class = ClassModel(name='Class 2')
        db.session.add(second_class)
        db.session.flush()
        extra = _add_student(second_class, 2001, 'Bilal Ahmed')
        first = db.session.get(StudentModel, seed['student_id'])

        _add_attendance(first, MONDAY, 'Present', seed['class_id'])
        _add_attendance(extra, MONDAY, 'Absent', second_class.id)
        db.session.commit()

        summary = build_school_attendance_summary(MONDAY, MONDAY)
        ids = {row['student'].id for row in summary['matrix_data']}
        assert first.id in ids
        assert extra.id in ids
        assert summary['kpi_stats']['total_students'] == len(ids)
        assert summary['class_obj'] is None
        assert summary['classes'][second_class.id].name == 'Class 2'


def test_attendance_summary_page_supports_whole_school(admin_client, app, seed):
    with app.app_context():
        _settings()
        _clear_holidays()
        student = db.session.get(StudentModel, seed['student_id'])
        _add_attendance(student, MONDAY, 'Present', seed['class_id'])
        db.session.commit()

    page = admin_client.get(
        f'/attendance/summary?class_id=all&start_date={MONDAY}&end_date={SUNDAY}')
    assert page.status_code == 200
    body = page.get_data(as_text=True)
    assert 'All School Students' in body
    assert 'Whole-School Attendance Summary' in body
    # Both students in the seed data are listed, across classes.
    assert 'Ali Khan' in body
    assert 'Sara Ali' in body
    assert 'Class 1' in body


def test_attendance_summary_page_still_defaults_to_a_class(admin_client, seed):
    response = admin_client.get('/attendance/summary')
    assert response.status_code == 302
    assert 'class_id=' in response.headers['Location']


# ==========================================
# EXPORTS
# ==========================================

def test_summary_exports_work_for_class_and_whole_school(admin_client, app, seed):
    with app.app_context():
        _settings()
        _clear_holidays()
        student = db.session.get(StudentModel, seed['student_id'])
        _add_attendance(student, MONDAY, 'Present', seed['class_id'])
        _add_attendance(student, TUESDAY, 'Late', seed['class_id'], late_minutes=17)
        db.session.commit()

    query = f'&start_date={MONDAY}&end_date={SUNDAY}'
    for class_id in (str(seed['class_id']), 'all'):
        pdf = admin_client.get(f'/attendance/summary/export/pdf?class_id={class_id}{query}')
        assert pdf.status_code == 200, class_id
        assert pdf.data.startswith(b'%PDF'), class_id
        assert 'attachment' in pdf.headers['Content-Disposition']

        xlsx = admin_client.get(f'/attendance/summary/export/excel?class_id={class_id}{query}')
        assert xlsx.status_code == 200, class_id
        # XLSX is a zip container: PK header.
        assert xlsx.data[:2] == b'PK', class_id

        csv_response = admin_client.get(f'/attendance/summary/export/csv?class_id={class_id}{query}')
        assert csv_response.status_code == 200, class_id
        text = csv_response.get_data(as_text=True)
        assert 'School' in text
        assert 'Late Minutes' in text
        assert 'Attendance days' in text or 'Period' in text
        if class_id == 'all':
            header_row = next(r for r in csv.reader(io.StringIO(text)) if r and r[0] == 'Sr')
            assert 'Class' in header_row
        # Late minutes are carried into the export.
        assert 'L (17m)' in text


def test_summary_exports_include_school_logo_when_configured(admin_client, app, seed):
    with app.app_context():
        settings = _settings()
        # Write a real (valid) PNG into the app static folder and point the setting at it.
        from PIL import Image as PILImage
        import os
        static_dir = os.path.join(app.root_path, 'static')
        logo_name = 'test_logo_attendance.png'
        logo_path = os.path.join(static_dir, logo_name)
        PILImage.new('RGB', (120, 60), (13, 110, 253)).save(logo_path)
        settings.logo_filename = logo_name
        student = db.session.get(StudentModel, seed['student_id'])
        _add_attendance(student, MONDAY, 'Present', seed['class_id'])
        db.session.commit()

        try:
            # The logo must actually be located and embedded by the exporter.
            from app.services.attendance_export import school_branding
            assert school_branding(app.root_path)['logo_path'] == logo_path

            pdf = admin_client.get(
                f'/attendance/summary/export/pdf?class_id={seed["class_id"]}'
                f'&start_date={MONDAY}&end_date={SUNDAY}')
            assert pdf.status_code == 200
            assert pdf.data.startswith(b'%PDF')
            assert b'/Image' in pdf.data

            xlsx = admin_client.get(
                f'/attendance/summary/export/excel?class_id={seed["class_id"]}'
                f'&start_date={MONDAY}&end_date={SUNDAY}')
            assert xlsx.status_code == 200
            assert xlsx.data[:2] == b'PK'
        finally:
            settings.logo_filename = ''
            db.session.commit()
            os.remove(logo_path)


def test_daily_class_exports_carry_late_minutes(admin_client, app, seed):
    with app.app_context():
        _settings()
        student = db.session.get(StudentModel, seed['student_id'])
        _add_attendance(student, MONDAY, 'Late', seed['class_id'], late_minutes=12)
        db.session.commit()

    csv_response = admin_client.get(
        f'/attendance/export/csv/{seed["class_id"]}?date={MONDAY}')
    assert csv_response.status_code == 200
    text = csv_response.get_data(as_text=True)
    assert 'Late Minutes' in text
    assert '12' in text

    xlsx = admin_client.get(f'/attendance/export/excel/{seed["class_id"]}?date={MONDAY}')
    assert xlsx.status_code == 200
    assert xlsx.data[:2] == b'PK'

    pdf = admin_client.get(f'/attendance/export/pdf/{seed["class_id"]}?date={MONDAY}')
    assert pdf.status_code == 200
    assert pdf.data.startswith(b'%PDF')


def _xlsx_data_rows(data):
    """Every non-empty spreadsheet row, as value tuples."""
    from openpyxl import load_workbook
    sheet = load_workbook(io.BytesIO(data)).active
    return [row for row in sheet.iter_rows(values_only=True)
            if row and any(cell is not None for cell in row)]


def _row_for(rows, name):
    return next(row for row in rows if name in row)


def _csv_rows(text):
    return list(csv.reader(io.StringIO(text)))


def _assert_late_columns(header, row, label):
    """Late Time and Late Minutes must hold the time and the minutes respectively.

    Guards against swapped positional arguments in the export call sites: the
    two maps are both {id: value} dicts, so a swap would silently put '08:47'
    into the minutes column rather than raising.
    """
    time_index = header.index('Late Time')
    minutes_index = header.index('Late Minutes')
    assert row[time_index] == LATE_TIME, \
        f'{label}: Late Time column holds {row[time_index]!r}, expected {LATE_TIME!r}'
    assert str(row[minutes_index]) == str(LATE_MINUTES), \
        f'{label}: Late Minutes column holds {row[minutes_index]!r}, expected {LATE_MINUTES!r}'


def test_daily_exports_place_late_time_and_late_minutes_in_correct_columns(
        admin_client, app, seed):
    """Regression guard for swapped late_time_map / late_minutes_map arguments."""
    with app.app_context():
        _settings()
        student = db.session.get(StudentModel, seed['student_id'])
        _add_attendance(student, MONDAY, 'Late', seed['class_id'], late_minutes=17)
        record = AttendanceModel.query.filter_by(
            target_type='student', target_id=student.id, date=MONDAY).first()
        record.late_time = LATE_TIME
        record.late_minutes = LATE_MINUTES
        db.session.commit()

    # --- XLSX ---
    rows = _xlsx_data_rows(
        admin_client.get(f'/attendance/export/excel/{seed["class_id"]}?date={MONDAY}').data)
    header = next(r for r in rows if 'Late Minutes' in r)
    _assert_late_columns(header, _row_for(rows, 'Ali Khan'), 'class xlsx')

    # --- CSV ---
    csv_rows = _csv_rows(admin_client.get(
        f'/attendance/export/csv/{seed["class_id"]}?date={MONDAY}').get_data(as_text=True))
    header = next(r for r in csv_rows if 'Late Minutes' in r)
    _assert_late_columns(header, _row_for(csv_rows, 'Ali Khan'), 'class csv')



def test_teacher_exports_include_late_minutes(admin_client, app):
    with app.app_context():
        _settings()
        teacher = TeacherModel(
            teacher_id_str='T-100', teacher_name='Miss Hina',
            qualification='MSc', is_active=True, monthly_salary=50000.0,
        )
        db.session.add(teacher)
        db.session.flush()
        db.session.add(AttendanceModel(
            target_type='teacher', target_id=teacher.id, date=MONDAY,
            status='Late', late_time='08:55', late_minutes=25,
        ))
        db.session.commit()
        teacher_id = teacher.id

    csv_response = admin_client.get(f'/attendance/teachers/export/csv?date={MONDAY}')
    assert csv_response.status_code == 200
    text = csv_response.get_data(as_text=True)
    assert 'Late Minutes' in text
    assert 'Miss Hina' in text
    assert '25' in text

    xlsx = admin_client.get(f'/attendance/teachers/export/excel?date={MONDAY}')
    assert xlsx.status_code == 200
    assert xlsx.data[:2] == b'PK'

    pdf = admin_client.get(f'/attendance/teachers/export/pdf?date={MONDAY}')
    assert pdf.status_code == 200
    assert pdf.data.startswith(b'%PDF')

    with app.app_context():
        assert db.session.get(TeacherModel, teacher_id) is not None


def test_teacher_summary_exports_include_late_minutes(admin_client, app):
    with app.app_context():
        _settings()
        teacher = TeacherModel(
            teacher_id_str='T-200', teacher_name='Mr Kamran',
            qualification='MEd', is_active=True, monthly_salary=60000.0,
        )
        db.session.add(teacher)
        db.session.flush()
        db.session.add(AttendanceModel(
            target_type='teacher', target_id=teacher.id, date=TUESDAY,
            status='Late', late_time='09:10', late_minutes=40,
        ))
        db.session.commit()

    base = f'/attendance/teachers/summary/export/{{fmt}}?start_date={MONDAY}&end_date={FRIDAY}'

    csv_response = admin_client.get(base.format(fmt='csv'))
    assert csv_response.status_code == 200
    text = csv_response.get_data(as_text=True)
    assert 'Late Minutes' in text
    assert 'Mr Kamran' in text
    assert '40' in text
    assert 'L (40m)' in text

    pdf = admin_client.get(base.format(fmt='pdf'))
    assert pdf.status_code == 200
    assert pdf.data.startswith(b'%PDF')

    xlsx = admin_client.get(base.format(fmt='excel'))
    assert xlsx.status_code == 200
    assert xlsx.data[:2] == b'PK'


def test_teacher_summary_page_shows_export_buttons(admin_client):
    page = admin_client.get(
        f'/attendance/teachers/summary?start_date={MONDAY}&end_date={FRIDAY}')
    assert page.status_code == 200
    body = page.get_data(as_text=True)
    assert '/attendance/teachers/summary/export/pdf' in body
    assert '/attendance/teachers/summary/export/csv' in body
    assert '/attendance/teachers/summary/export/excel' in body


def test_summary_csv_parses_with_expected_header(admin_client, app, seed):
    with app.app_context():
        _settings()
        _clear_holidays()
        student = db.session.get(StudentModel, seed['student_id'])
        _add_attendance(student, MONDAY, 'Present', seed['class_id'])
        db.session.commit()

    response = admin_client.get(
        f'/attendance/summary/export/csv?class_id={seed["class_id"]}'
        f'&start_date={MONDAY}&end_date={SUNDAY}')
    rows = list(csv.reader(io.StringIO(response.get_data(as_text=True))))
    header_row = next(r for r in rows if r and r[0] == 'Sr')
    assert header_row[:2] == ['Sr', 'Roll No']
    assert 'Late Minutes' in header_row
    assert 'Percentage' in header_row
    # The working Monday is a column, the weekend is not.
    assert '21/09/2026' in header_row
    assert '26/09/2026' not in header_row
    assert '27/09/2026' not in header_row
