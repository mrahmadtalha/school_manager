"""Student / Teacher import: mandatory fields, duplicates, auto numbering, templates."""

import io
from datetime import date

import openpyxl

from app.database import db
from app.models import ClassModel, StudentEnrollment, StudentModel, TeacherModel
from app.services.roll_numbers import FIRST_ROLL_NUMBER

STUDENT_HEADER = ('Student Name,Father Name,Class,Guardian Phone,Address,'
                  'Roll Number,Section,Admission No,Date of Birth\n')
TEACHER_HEADER = ('Teacher Name,Joining Date,Qualification,Teacher ID,CNIC,'
                  'Contact Number,Assigned Classes,Monthly Salary\n')


def _post(client, url, csv_text, filename='import.csv'):
    return client.post(url, data={
        'import_file': (io.BytesIO(csv_text.encode('utf-8')), filename),
    }, content_type='multipart/form-data', follow_redirects=True)


def _students(client, body):
    response = _post(client, '/students/import', STUDENT_HEADER + body)
    assert response.status_code == 200
    return response.get_data(as_text=True)


def _teachers(client, body):
    response = _post(client, '/teachers/import', TEACHER_HEADER + body)
    assert response.status_code == 200
    return response.get_data(as_text=True)


# -- students: mandatory fields ----------------------------------------------

def test_student_missing_mandatory_value_is_skipped(admin_client, app, seed):
    page = _students(admin_client,
                     'No Phone,Father,Class 1,,Street,,,,\n'
                     'No Address,Father,Class 1,03001112222,,,,,\n'
                     'No Class,Father,,03001112222,Street,,,,\n'
                     'Good Student,Father,Class 1,03001112222,Street,,,,\n')
    assert 'Missing required value(s)' in page
    with app.app_context():
        assert StudentModel.query.filter_by(student_name='No Phone').count() == 0
        assert StudentModel.query.filter_by(student_name='No Address').count() == 0
        assert StudentModel.query.filter_by(student_name='No Class').count() == 0
        assert StudentModel.query.filter_by(student_name='Good Student').count() == 1


def test_student_file_without_mandatory_column_is_rejected(admin_client, app, seed):
    response = _post(admin_client, '/students/import',
                     'Student Name,Class\nAli,Class 1\n')
    assert b'missing the required column' in response.data
    with app.app_context():
        assert StudentModel.query.count() == 2          # only the two seeded students


# -- students: duplicates ------------------------------------------------------

def test_student_duplicates_are_blocked(admin_client, app, seed):
    page = _students(admin_client,
                     'Roll Clash,Father A,Class 1,03001112222,Street,1001,,,\n'
                     'Ali Khan,Imran Khan,Class 1,03009998888,Street,,,,\n'
                     'Adm One,Father B,Class 1,03001112222,Street,,,ADM-1,\n'
                     'Adm Two,Father C,Class 1,03001112222,Street,,,adm-1,\n')
    assert 'already registered' in page
    with app.app_context():
        assert StudentModel.query.filter_by(student_name='Roll Clash').count() == 0
        assert StudentModel.query.filter_by(student_name='Ali Khan').count() == 1
        assert StudentModel.query.filter_by(student_name='Adm One').count() == 1
        assert StudentModel.query.filter_by(student_name='Adm Two').count() == 0


def test_student_duplicate_inside_the_file_is_blocked(admin_client, app, seed):
    _students(admin_client,
              'Twin Roll A,F,Class 1,03001112222,Street,1050,,,\n'
              'Twin Roll B,G,Class 1,03001112222,Street,1050,,,\n'
              'Same Person,Dad,Class 1,03001112222,Street,,,,\n'
              'same  person,dad,Class 1,03001112222,Street,,,,\n')
    with app.app_context():
        assert StudentModel.query.filter_by(student_name='Twin Roll A').count() == 1
        assert StudentModel.query.filter_by(student_name='Twin Roll B').count() == 0
        assert StudentModel.query.filter(
            StudentModel.student_name.ilike('same person')).count() == 1


def test_importing_the_same_file_twice_adds_nothing_new(admin_client, app, seed):
    body = ('Repeat One,Dad One,Class 1,03001112222,Street,,,,\n'
            'Repeat Two,Dad Two,Class 1,03001112222,Street,,,,\n')
    _students(admin_client, body)
    _students(admin_client, body)
    with app.app_context():
        assert StudentModel.query.filter(
            StudentModel.student_name.like('Repeat %')).count() == 2


# -- students: automatic roll numbers ----------------------------------------

def test_blank_roll_numbers_are_assigned_automatically(admin_client, app, seed):
    _students(admin_client,
              'Auto One,Dad,Class 1,03001112222,Street,,,,\n'
              'Explicit,Dad,Class 1,03001112222,Street,1003,,,\n'
              'Auto Two,Dad,Class 1,03001112222,Street,,,,\n')
    with app.app_context():
        rolls = {s.student_name: s.roll_number for s in StudentModel.query.all()}
        assert rolls['Explicit'] == FIRST_ROLL_NUMBER + 2          # 1003 as typed
        assert rolls['Auto One'] == FIRST_ROLL_NUMBER + 3          # 1004 (1003 is taken)
        assert rolls['Auto Two'] == FIRST_ROLL_NUMBER + 4
        assert len(set(rolls.values())) == len(rolls)


def test_first_auto_roll_in_empty_class_is_first_number(admin_client, app, seed):
    with app.app_context():
        db.session.add(ClassModel(name='Class 2'))
        db.session.commit()
    _post(admin_client, '/students/import',
          STUDENT_HEADER + 'Fresh,Dad,Class 2,03001112222,Street,,,,\n')
    with app.app_context():
        assert StudentModel.query.filter_by(student_name='Fresh').one().roll_number == FIRST_ROLL_NUMBER


def test_imported_student_gets_enrollment_and_fields(admin_client, app, seed):
    _students(admin_client,
              'Full Row,Dad,Class 1,03001112222,Street,,A,ADM-9,2015-04-20\n')
    with app.app_context():
        student = StudentModel.query.filter_by(student_name='Full Row').one()
        assert student.section_id == seed['section_id']
        assert student.admission_number == 'ADM-9'
        assert student.date_of_birth == date(2015, 4, 20)
        assert StudentEnrollment.query.filter_by(student_id=student.id).count() == 1


def test_student_unknown_class_or_section_is_skipped(admin_client, app, seed):
    page = _students(admin_client,
                     'Bad Class,Dad,Class 99,03001112222,Street,,,,\n'
                     'Bad Section,Dad,Class 1,03001112222,Street,,Z,,\n'
                     'Bad Date,Dad,Class 1,03001112222,Street,,,,not-a-date\n')
    assert 'does not exist' in page
    with app.app_context():
        for name in ('Bad Class', 'Bad Section', 'Bad Date'):
            assert StudentModel.query.filter_by(student_name=name).count() == 0


# -- teachers ------------------------------------------------------------------

def test_teacher_missing_mandatory_value_is_skipped(admin_client, app, seed):
    page = _teachers(admin_client,
                     'No Qualification,2024-01-15,,,,,,\n'
                     'No Date,,M.A,,,,,\n'
                     ',2024-01-15,M.A,,,,,\n'
                     'Good Teacher,2024-01-15,M.A English,,,,,40000\n')
    assert 'Missing required value(s)' in page
    with app.app_context():
        assert TeacherModel.query.count() == 1
        teacher = TeacherModel.query.one()
        assert teacher.teacher_name == 'Good Teacher'
        assert teacher.teacher_id_str == 'T001'          # assigned automatically
        assert teacher.monthly_salary == 40000.0


def test_teacher_duplicates_are_blocked(admin_client, app, seed):
    _teachers(admin_client,
              'First,2024-01-15,M.A,T010,36302-1234567-1,03001112222,,30000\n')
    page = _teachers(admin_client,
                     'Same ID,2024-02-01,B.Ed,T010,,,,\n'
                     'Same CNIC,2024-02-01,B.Ed,T011,3630212345671,,,\n'
                     'First,2024-03-01,B.Sc,,,03001112222,,\n'
                     'Fresh Teacher,2024-02-01,B.Ed,,,,,\n')
    assert 'already registered' in page
    with app.app_context():
        names = sorted(t.teacher_name for t in TeacherModel.query.all())
        assert names == ['First', 'Fresh Teacher']
        fresh = TeacherModel.query.filter_by(teacher_name='Fresh Teacher').one()
        assert fresh.teacher_id_str == 'T011'            # T010 is taken -> next free


def test_teacher_duplicates_inside_the_file_are_blocked(admin_client, app, seed):
    _teachers(admin_client,
              'Auto,2024-01-15,M.A,,,,,\n'
              'Explicit,2024-01-15,M.A,T001,,,,\n'
              'Repeat ID,2024-01-15,M.A,T001,,,,\n')
    with app.app_context():
        ids = {t.teacher_name: t.teacher_id_str for t in TeacherModel.query.all()}
        assert ids['Explicit'] == 'T001'
        assert ids['Auto'] == 'T002'                      # T001 reserved by a later row
        assert 'Repeat ID' not in ids


def test_teacher_unknown_class_is_skipped_known_class_saved(admin_client, app, seed):
    _teachers(admin_client,
              'Wrong Class,2024-01-15,M.A,,,,Class 77,\n'
              'Right Class,2024-01-15,M.A,,,,class 1,\n')
    with app.app_context():
        assert TeacherModel.query.filter_by(teacher_name='Wrong Class').count() == 0
        right = TeacherModel.query.filter_by(teacher_name='Right Class').one()
        assert right.assigned_class == 'Class 1'


# -- templates -----------------------------------------------------------------

def _template_sheet(client, url):
    response = client.get(url)
    assert response.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(response.data))
    return workbook, workbook.worksheets[0]


def test_student_template_has_all_fields_example_and_markers(admin_client):
    workbook, sheet = _template_sheet(admin_client, '/students/template/excel')
    headers = [cell.value for cell in sheet[1]]
    for mandatory in ('Student Name *', 'Father Name *', 'Class *',
                      'Guardian Phone *', 'Address *'):
        assert mandatory in headers
    for optional in ('Roll Number', 'Section', 'Gender', 'Date of Birth', 'Admission No',
                     'Admission Date', 'Sponsor Type', 'Sponsor CNIC', 'Monthly Fee',
                     'Discount Type', 'Discount Value'):
        assert optional in headers
    assert sheet['A2'].value == 'Ali Khan'                    # one example row
    assert sheet.max_row >= 2 and sheet['A3'].value is None
    assert 'Instructions' in workbook.sheetnames


def test_teacher_template_has_all_fields_example_and_markers(admin_client):
    workbook, sheet = _template_sheet(admin_client, '/teachers/template/excel')
    headers = [cell.value for cell in sheet[1]]
    for mandatory in ('Teacher Name *', 'Joining Date *', 'Qualification *'):
        assert mandatory in headers
    for optional in ('Teacher ID', 'Designation', 'Gender', 'Salary Type', 'Monthly Salary',
                     'Hourly Rate', 'Date of Birth', 'CNIC', 'Address', 'Contact Number',
                     'Emergency Contact Number', 'Assigned Classes', 'Assigned Subjects',
                     'Previous Experience (Years)', 'Previous Salary'):
        assert optional in headers
    assert sheet['A2'].value == 'Ayesha Malik'
    assert 'Instructions' in workbook.sheetnames


def test_untouched_template_example_row_is_ignored_on_import(admin_client, app, seed):
    response = admin_client.get('/students/template/excel')
    result = admin_client.post('/students/import', data={
        'import_file': (io.BytesIO(response.data), 'students_import_template.xlsx'),
    }, content_type='multipart/form-data', follow_redirects=True)
    assert result.status_code == 200
    with app.app_context():
        assert StudentModel.query.filter_by(student_name='Ali Khan').count() == 1   # seed only
